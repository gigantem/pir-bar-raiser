#!/usr/bin/env python3
# SPDX-License-Identifier: GPL-3.0-or-later
# Copyright (C) 2026 Mike Gigante <mike.gigante@gmail.com>
"""render_pir_docx.py: the PIR template, and a finished PIR, from one layout.

One renderer, two inputs. Without `--packet` it produces the blank template: every
incident-specific field a yellow bracketed placeholder, a grey guidance box under each heading
(the catalog's `template_guidance`), and two appendices. Appendix A states the expectation and
the common mistakes for each section from catalogs/pir-section-expectations.yaml and shows the
worked example rendered from the example packet itself (examples/good/), so the appendix cannot
disagree with what the validator accepts. Appendix B is catalogs/pir-principles.yaml. With
`--packet` it renders that packet into the same layout, section for section, with no guidance
boxes and no appendices unless asked for, so the published document and the template are the
same shape.

Both inputs render to `.docx` (python-docx) or to markdown (`--format md`). `--check` compares
the document's outline (tools/docx_outline.py: every paragraph with its style, in order) with an
expected outline, and the markdown with expected markdown, so a rendering can be compared with
one written by another tool.

Nothing is invented: an absent packet section renders as "not yet provided", never as a
plausible value.
"""

from __future__ import annotations

import argparse
import difflib
import io
import sys
import tempfile
import zipfile
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import docx_outline
import pir_interview
import validate_pir

REPO_ROOT = Path(__file__).resolve().parent.parent
PRINCIPLES_PATH = REPO_ROOT / "catalogs" / "pir-principles.yaml"
DEFAULT_EXAMPLE_PACKET = REPO_ROOT / "examples" / "good" / "order-capture-pool-exhaustion.yaml"
TEMPLATE_VERSION = "v0.2"
FONT = "Arial"
NAVY = RGBColor(0x1F, 0x2A, 0x44)
ACCENT = RGBColor(0x0B, 0x5F, 0xFF)
GREY_TXT = RGBColor(0x59, 0x59, 0x59)
PLACEHOLDER_TXT = RGBColor(0x7F, 0x60, 0x00)
EXAMPLE_TXT = RGBColor(0x1E, 0x3A, 0x2A)
EXAMPLE_LABEL = RGBColor(0x1E, 0x5A, 0x38)
PH_FILL, GUIDE_FILL, HEAD_FILL, EX_FILL = "FFF2CC", "F2F2F2", "DCE6F7", "E8F3EC"
CONTENT_WIDTH_CM = 17.0
NOT_PROVIDED = "not yet provided"
#: Fixed so two renders of the same input produce the same core properties.
FIXED_TIMESTAMP = datetime(2026, 1, 1, tzinfo=UTC)
#: Every entry of a written .docx carries this stamp, the earliest a zip can record, so a document
#: says nothing about when, or in which timezone, it was written, and the same packet always
#: renders to the same bytes.
ZIP_DATE_TIME = (1980, 1, 1, 0, 0, 0)
BODY_ORDER = (
    "summary",
    "timeline",
    "metrics",
    "customer_impact",
    "response_analysis",
    "five_whys",
    "lessons",
    "corrective_actions",
    "action_items",
    "review",
)
BODY_TITLE_OVERRIDES = {"five_whys": "Five Whys (root cause analysis)"}
CHAIN_TITLES = {
    "failure": "the failure",
    "detection": "the detection gap",
    "resolution": "the slow resolution",
    "prior_actions": "why earlier actions did not prevent this",
}
PHASE_LABELS = {
    "precursor": "Precursor",
    "impact_start": "Impact start",
    "detection": "Detection",
    "response": "Response",
    "diagnosis": "Diagnosis",
    "remediation": "Remediation",
    "mitigated": "Mitigated",
    "resolved": "Resolved",
    "comms": "Comms",
}
IMPROVES_LABELS = {"prevention": "Prevention", "ttd": "TTD", "ttk": "TTK", "ttr": "TTR"}
STAGE_LABELS = {
    "team": "Team review",
    "go": "GO review",
    "exec_go": "Executive GO review",
    "stakeholder_alignment": "Customer / stakeholder alignment",
}


@dataclass(frozen=True)
class Placeholder:
    """An incident-specific field the author fills in: rendered yellow, bracketed, italic."""

    text: str


Cell = str | Placeholder
Row = list[Cell]


def ph(text: str) -> Placeholder:
    return Placeholder(text)


# ================================================================================ the model


def template_model() -> dict[str, Any]:
    """Every section of the template with placeholders where the packet's values go."""
    whys = [(f"Why {i}", ph("")) for i in range(1, 6)]
    chains = []
    for letter, kind in zip("ABCD", ("failure", "detection", "resolution", "prior_actions"), strict=True):
        title = CHAIN_TITLES[kind] + (" (delete if not a repeat)" if kind == "prior_actions" else "")
        chains.append(
            {
                "letter": letter,
                "title": title,
                "problem": ph("One sentence, stated as an observed fact"),
                "whys": whys,
                "root_cause": ph("The process, design or tooling gap at the end of the chain. Not a person."),
            }
        )
    return {
        "incident_id": ph("Incident ID"),
        "title": ph("Short incident title"),
        "classification": ph("General"),
        "control": [
            ("Incident ID", ph("e.g. INC-2026-0412")),
            ("Incident title", ph("Service or application: what customers experienced, in one line")),
            ("Severity", ph("Sev-1 / Sev-2 / Sev-3")),
            ("Service / application", ph("Primary service affected, plus dependent services if any")),
            (
                "Impact start",
                ph("YYYY-MM-DD HH:MM TZ: when customers first experienced impact (not when we noticed)"),
            ),
            ("Detected at", ph("YYYY-MM-DD HH:MM TZ: first signal of any kind, a customer report included")),
            ("Cause confirmed at", ph("YYYY-MM-DD HH:MM TZ: the cause proven with evidence")),
            ("Mitigated at", ph("YYYY-MM-DD HH:MM TZ: customer impact stopped or materially reduced")),
            ("Resolved at", ph("YYYY-MM-DD HH:MM TZ: service fully restored and incident closed")),
            ("Customer impact duration", ph("Mitigated at minus impact start, e.g. 74 min")),
            ("PIR owner (Product Owner)", ph("Role, team and name. The PO owns the PIR and presents it.")),
            ("Incident Manager", ph("Role and team")),
            ("Contributors / SMEs", ph("Teams and roles consulted, including third parties")),
            ("Repeat incident", ph("Yes, with the prior PIR reference / No")),
            ("PIR status", ph("Draft / In team review / In GO review / Published")),
            ("Draft due", ph("Impact-mitigated date + 2 to 3 business days")),
            ("Publish due", ph("Impact-mitigated date + 1 to 2 weeks")),
            (
                "Related records",
                ph("Change IDs, problem tickets, prior PIRs for repeat incidents, war-room link"),
            ),
            ("Classification", ph("e.g. General / Internal / Confidential")),
        ],
        "summary": ph(
            "Application experienced metric at date time. What happened, rough timeline. Customer impact. "
            "Current state and whether the fix is permanent or interim."
        ),
        "key_timings": [
            ("Time to detect (TTD)", "Impact start to detected at", ph("e.g. 9 min")),
            ("Time to know (TTK)", "Detected at to cause confirmed with evidence", ph("e.g. 44 min")),
            ("Time to remediate (TTR)", "Cause confirmed to mitigated at", ph("e.g. 21 min")),
            ("Customer impact duration", "Impact start to mitigated at", ph("e.g. 74 min")),
            ("Time to resolve", "Impact start to resolved at", ph("e.g. 95 min")),
        ],
        "timeline": [
            [
                ph("YYYY-MM-DD HH:MM TZ"),
                ph("Triggering change, deploy, or external event (if known)"),
                ph("Change ID, deploy log"),
                "Precursor",
            ],
            [
                ph("HH:MM"),
                ph("Customer impact begins (first affected request / transaction)"),
                ph("Graph, log query"),
                "Impact start",
            ],
            [
                ph("HH:MM"),
                ph("First signal: alert fired / customer report / ops noticed"),
                ph("Alert name, ticket"),
                "Detection",
            ],
            [
                ph("HH:MM"),
                ph("Incident declared at Sev-N; IM engaged; war room opened"),
                ph("Decision log"),
                "Response",
            ],
            [
                ph("HH:MM"),
                ph("Hypothesis raised / proven / rejected, with the evidence used"),
                ph("Dashboard, query"),
                "Diagnosis",
            ],
            [ph("HH:MM"), ph("Cause confirmed"), ph(""), "Diagnosis"],
            [
                ph("HH:MM"),
                ph("Mitigation decided and started (e.g. rollback, failover, scale-out)"),
                ph(""),
                "Remediation",
            ],
            [ph("HH:MM"), ph("Customer impact stops or materially reduces"), ph("Graph"), "Mitigated"],
            [
                ph("HH:MM"),
                ph("Service fully restored; backlog cleared; incident resolved"),
                ph(""),
                "Resolved",
            ],
            [ph("HH:MM"), ph("Final stakeholder / customer communication sent"), ph("Comms record"), "Comms"],
        ],
        "figures": [
            {
                "label": "Figure 1.",
                "insert": ph(
                    "Insert Figure 1 here: the primary customer-facing metric over the incident window, with impact "
                    "start, detection, mitigation and resolution annotated, and the same window from a normal day for comparison."
                ),
                "caption": ph("Metric name, service, time range and timezone, data source"),
            },
            {
                "label": "Figure 2.",
                "insert": ph(
                    "Insert Figure 2 here: the diagnostic metric that confirmed the cause (e.g. saturation, error class, "
                    "queue depth) on the same time range."
                ),
                "caption": ph("Caption"),
            },
            {
                "label": "Figure 3.",
                "insert": ph("Insert a diagram here if the failure path crosses more than two components."),
                "caption": ph("Caption"),
            },
        ],
        "metric_summary": [
            [
                ph("e.g. 5xx error rate (per min)"),
                ph("Customer-facing failure"),
                ph("0.2%"),
                ph("Peak 11%"),
                ph("Dashboard link"),
            ],
            [
                ph("e.g. Order submissions succeeded (per min)"),
                ph("Business outcome"),
                ph(""),
                ph(""),
                ph(""),
            ],
            [ph("e.g. DB pool connections in use"), ph("Diagnostic / saturation"), ph(""), ph(""), ph("")],
        ],
        "impact": [
            ("Who was affected (segments, regions, channels, partners)?", ph("")),
            ("What did they experience?", ph("In plain language: what the customer saw or could not do")),
            (
                "How many were affected? How do we know?",
                ph("Count, percentage of normal volume, and the method used to count"),
            ),
            ("For how long?", ph("Impact window, including any partial-degradation tail")),
            (
                "Business metrics affected",
                ph("e.g. failed orders, revenue at risk, delayed activations, support contacts"),
            ),
            (
                "Commitments breached (SLA, SLO, regulatory)?",
                ph("Which, by how much, and any credits or notifications owed"),
            ),
            ("Customers / partners contacted?", ph("Who, when, by whom (role), and what they told us")),
            ("What was not affected?", ph("")),
            (
                "Any data loss, duplication or corruption?",
                ph('State explicitly, including "none" if verified and how'),
            ),
        ],
        "went_well": [
            ph("Specific things that shortened or contained the incident, and why they worked"),
            ph(""),
        ],
        "missed": [
            [
                "Detection (time to detect)",
                ph(
                    "e.g. no alert on the customer-facing metric; threshold too slow; alert went to the wrong team"
                ),
                ph("e.g. 9 min with no signal at all"),
            ],
            [
                "Diagnosis (time to know)",
                ph("e.g. diagnostic metric not on the dashboard; wrong first hypothesis; missing access"),
                ph(""),
            ],
            [
                "Remediation (time to remediate)",
                ph("e.g. rollback not automated; runbook stale; change approval delay"),
                ph(""),
            ],
            [
                "Coordination and comms",
                ph("e.g. late escalation, unclear roles, stakeholder surprise"),
                ph(""),
            ],
            [
                "Prevention",
                ph("e.g. change not tested at load; known risk not acted on; prior PIR action not completed"),
                ph(""),
            ],
        ],
        "channels": [
            ph("The failure itself"),
            ph("Why it was not detected sooner"),
            ph("Why it took as long as it did to fix"),
            ph("If this has happened before: why previous actions did not prevent it"),
        ],
        "chains": chains,
        "lessons": [
            ["LL-1", ph(""), ph("RC-A"), ph("CA-1")],
            ["LL-2", ph(""), ph("RC-B"), ph("CA-2")],
            ["LL-3", ph(""), ph("RC-C"), ph("CA-3")],
        ],
        "corrective_actions": [
            ["CA-1", ph(""), ph("LL-1"), ph("Prevention / TTD / TTK / TTR"), ph("AI-1, AI-2")],
            ["CA-2", ph(""), ph("LL-2"), ph(""), ph("")],
            ["CA-3", ph(""), ph("LL-3"), ph(""), ph("")],
        ],
        "action_items": [
            [
                "AI-1",
                ph(""),
                ph("CA-1"),
                ph("Person or team who agreed"),
                ph("YYYY-MM-DD"),
                ph("Open"),
                ph(""),
            ],
            ["AI-2", ph(""), ph(""), ph(""), ph(""), ph(""), ph("")],
            ["AI-3", ph(""), ph(""), ph(""), ph(""), ph(""), ph("")],
        ],
        "review_stages": [
            ["Team review", ph(""), ph(""), ph("")],
            ["GO review", ph(""), ph(""), ph("")],
            ["Executive GO review (if warranted)", ph(""), ph(""), ph("")],
            ["Customer / stakeholder alignment", ph(""), ph(""), ph("")],
        ],
        "open_questions": [[ph(""), ph(""), ph(""), ph("")]],
    }


def _text(value: Any, absent: str = NOT_PROVIDED) -> str:
    if value is None or value == "" or value == []:
        return absent
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, list):
        return ", ".join(_text(v) for v in value)
    return str(value)


def _time_cell(value: Any, previous_date: str | None) -> tuple[str, str | None]:
    """Date and time with offset on the first row and whenever the date changes; time only otherwise."""
    parsed = validate_pir.parse_timestamp(value)
    if parsed is None:
        return _text(value), previous_date
    date = parsed.date().isoformat()
    offset = parsed.strftime("%z")
    offset = f"{offset[:3]}:{offset[3:]}" if offset else ""
    if date != previous_date:
        return f"{date} {parsed.strftime('%H:%M')} {offset}".strip(), date
    return parsed.strftime("%H:%M"), date


def _minutes(value: Any) -> str:
    return f"{value} min" if isinstance(value, int) else _text(value)


def packet_model(packet: dict[str, Any]) -> dict[str, Any]:
    """The same shape as template_model(), with the packet's values; absent fields say so."""
    control = validate_pir.mapping(packet.get("document_control"))
    stamps = validate_pir.mapping(control.get("timestamps"))
    parsed = validate_pir.lifecycle_timestamps(packet)
    timings = packet.get("key_timings") if isinstance(packet.get("key_timings"), dict) else None
    if timings is None and parsed is not None:
        timings = validate_pir.derive_key_timings(parsed)
    timings = timings or {}

    def stamp(key: str) -> str:
        value = validate_pir.parse_timestamp(stamps.get(key))
        if value is None:
            return _text(stamps.get(key))
        offset = value.strftime("%z")
        return f"{value.strftime('%Y-%m-%d %H:%M')} {offset[:3]}:{offset[3:]}"

    status_labels = {
        "draft": "Draft",
        "in_team_review": "In team review",
        "in_go_review": "In GO review",
        "published": "Published",
    }
    repeat = control.get("repeat_incident")
    repeat_text = NOT_PROVIDED if repeat is None else ("Yes" if repeat else "No")

    timeline_rows: list[Row] = []
    previous_date: str | None = None
    for row in packet.get("timeline") or []:
        if not isinstance(row, dict):
            continue
        when, previous_date = _time_cell(row.get("at"), previous_date)
        timeline_rows.append(
            [
                when,
                _text(row.get("event")),
                _text(row.get("evidence")),
                PHASE_LABELS.get(str(row.get("phase")), _text(row.get("phase"))),
            ]
        )

    metrics = validate_pir.mapping(packet.get("metrics"))
    figures = []
    for index, figure in enumerate(metrics.get("figures") or [], start=1):
        if not isinstance(figure, dict):
            continue
        flags = []
        flags.append("annotated" if figure.get("annotated") else "not annotated")
        flags.append("baseline compared" if figure.get("baseline_compared") else "no baseline")
        figures.append(
            {
                "label": f"Figure {index} ({figure.get('id', '?')}).",
                "insert": None,
                "caption": f"{_text(figure.get('caption'))} Source: {_text(figure.get('source'))}. Window: {_text(figure.get('time_range'))}. {', '.join(flags).capitalize()}.",
            }
        )
    metric_rows = [
        [
            _text(m.get("metric")),
            _text(m.get("meaning")),
            _text(m.get("normal_baseline")),
            _text(m.get("during_incident")),
            _text(m.get("source")),
        ]
        for m in metrics.get("summary") or []
        if isinstance(m, dict)
    ]

    impact = validate_pir.mapping(packet.get("customer_impact"))
    impact_rows = [
        ("Who was affected (segments, regions, channels, partners)?", _text(impact.get("who"))),
        ("What did they experience?", _text(impact.get("experienced"))),
        (
            "How many were affected? How do we know?",
            f"{_text(impact.get('how_many'))} Counted by: {_text(impact.get('how_counted'))}",
        ),
        ("For how long?", _text(impact.get("duration"))),
        ("Business metrics affected", _text(impact.get("business_metrics"))),
        ("Commitments breached (SLA, SLO, regulatory)?", _text(impact.get("commitments"))),
        ("Customers / partners contacted?", _text(impact.get("contacted"))),
        ("What was not affected?", _text(impact.get("not_affected"))),
        ("Any data loss, duplication or corruption?", _text(impact.get("data_integrity"))),
    ]

    analysis = validate_pir.mapping(packet.get("response_analysis"))
    area_labels = {
        "detection": "Detection (time to detect)",
        "diagnosis": "Diagnosis (time to know)",
        "remediation": "Remediation (time to remediate)",
        "coordination": "Coordination and comms",
        "prevention": "Prevention",
    }
    missed_rows = [
        [
            area_labels.get(str(m.get("area")), _text(m.get("area"))),
            _text(m.get("gap")),
            _text(m.get("effect")),
        ]
        for m in analysis.get("missed") or []
        if isinstance(m, dict)
    ]
    channels = [
        f"{c.get('id', '?')}: {_text(c.get('statement'))}"
        for c in analysis.get("channels_of_inquiry") or []
        if isinstance(c, dict)
    ]

    chains = []
    for chain in packet.get("five_whys") or []:
        if not isinstance(chain, dict):
            continue
        chain_id = str(chain.get("id", "RC-?"))
        letter = chain_id.rpartition("-")[2] or "?"
        kind = str(chain.get("kind", ""))
        chains.append(
            {
                "letter": letter,
                "title": f"{CHAIN_TITLES.get(kind, kind)} ({chain.get('channel', 'no channel')})",
                "problem": _text(chain.get("problem_statement")),
                "whys": [
                    (
                        f"Why {i}: {_text(w.get('why'))}",
                        f"{_text(w.get('because'))} Evidence: {_text(w.get('evidence'))}",
                    )
                    for i, w in enumerate(chain.get("whys") or [], start=1)
                    if isinstance(w, dict)
                ],
                "root_cause": _text(chain.get("root_cause")),
            }
        )

    corrective_actions = [c for c in packet.get("corrective_actions") or [] if isinstance(c, dict)]
    action_items = [a for a in packet.get("action_items") or [] if isinstance(a, dict)]
    lessons_rows = []
    for lesson in packet.get("lessons") or []:
        if not isinstance(lesson, dict):
            continue
        addressed_by = [
            str(c.get("id")) for c in corrective_actions if lesson.get("id") in (c.get("addresses") or [])
        ]
        lessons_rows.append(
            [
                _text(lesson.get("id")),
                _text(lesson.get("statement")),
                _text(lesson.get("root_cause")),
                _text(addressed_by, "none yet"),
            ]
        )
    ca_rows = []
    for ca in corrective_actions:
        items = [str(a.get("id")) for a in action_items if a.get("corrective_action") == ca.get("id")]
        improves = [IMPROVES_LABELS.get(str(x), str(x)) for x in ca.get("improves") or []]
        ca_rows.append(
            [
                _text(ca.get("id")),
                _text(ca.get("statement")),
                _text(ca.get("addresses")),
                _text(improves),
                _text(items, "none yet"),
            ]
        )
    ai_rows = []
    for ai in action_items:
        owner = _text(ai.get("owner"))
        if ai.get("owner_agreed") is False:
            owner += " (not yet agreed)"
        ai_rows.append(
            [
                _text(ai.get("id")),
                _text(ai.get("statement")),
                _text(ai.get("corrective_action")),
                owner,
                _text(ai.get("due_date")),
                _text(ai.get("status")).replace("_", " "),
                _text(ai.get("prevents_recurrence")),
            ]
        )

    review = validate_pir.mapping(packet.get("review"))
    stage_rows = [
        [
            STAGE_LABELS.get(str(s.get("stage")), _text(s.get("stage"))),
            _text(s.get("reviewer")),
            _text(s.get("date")),
            _text(s.get("outcome")),
        ]
        for s in review.get("stages") or []
        if isinstance(s, dict)
    ]
    question_rows = [
        [
            f"{q.get('id', '?')}: {_text(q.get('question'))}",
            _text(q.get("owner")),
            _text(q.get("answer_due")),
            _text(q.get("answer"), "open"),
        ]
        for q in review.get("open_questions") or []
        if isinstance(q, dict)
    ]

    return {
        "incident_id": _text(control.get("incident_id")),
        "title": _text(control.get("title")),
        "classification": _text(control.get("classification"), "General"),
        "control": [
            ("Incident ID", _text(control.get("incident_id"))),
            ("Incident title", _text(control.get("title"))),
            ("Severity", _text(control.get("severity"))),
            ("Service / application", _text(control.get("services"))),
            ("Impact start", stamp("impact_start")),
            ("Detected at", stamp("detected_at")),
            ("Cause confirmed at", stamp("cause_confirmed_at")),
            ("Mitigated at", stamp("mitigated_at")),
            ("Resolved at", stamp("resolved_at")),
            ("Customer impact duration", _minutes(timings.get("customer_impact_minutes"))),
            ("PIR owner (Product Owner)", _text(control.get("pir_owner"))),
            ("Incident Manager", _text(control.get("incident_manager"))),
            ("Contributors / SMEs", _text(control.get("contributors"))),
            ("Repeat incident", repeat_text),
            ("PIR status", status_labels.get(str(packet.get("status")), _text(packet.get("status")))),
            ("Draft due", _text(control.get("draft_due"))),
            ("Publish due", _text(control.get("publish_due"))),
            ("Related records", _text(control.get("related_records"))),
            ("Classification", _text(control.get("classification"))),
        ],
        "summary": _text(packet.get("summary")),
        "key_timings": [
            (
                "Time to detect (TTD)",
                "Impact start to detected at",
                _minutes(timings.get("time_to_detect_minutes")),
            ),
            (
                "Time to know (TTK)",
                "Detected at to cause confirmed with evidence",
                _minutes(timings.get("time_to_know_minutes")),
            ),
            (
                "Time to remediate (TTR)",
                "Cause confirmed to mitigated at",
                _minutes(timings.get("time_to_remediate_minutes")),
            ),
            (
                "Customer impact duration",
                "Impact start to mitigated at",
                _minutes(timings.get("customer_impact_minutes")),
            ),
            (
                "Time to resolve",
                "Impact start to resolved at",
                _minutes(timings.get("time_to_resolve_minutes")),
            ),
        ],
        "timeline": timeline_rows,
        "figures": figures,
        "metric_summary": metric_rows,
        "impact": impact_rows,
        "went_well": [_text(x) for x in analysis.get("went_well") or []],
        "missed": missed_rows,
        "channels": channels,
        "chains": chains,
        "lessons": lessons_rows,
        "corrective_actions": ca_rows,
        "action_items": ai_rows,
        "review_stages": stage_rows,
        "open_questions": question_rows,
    }


# ============================================================================ the appendix


def example_lines(section_key: str, model: dict[str, Any]) -> list[tuple[str, str]]:
    """(label, text) lines for Appendix A's example of one section, rendered from the example packet."""
    if section_key == "summary":
        timings = "; ".join(f"{measure} {value}" for measure, _, value in model["key_timings"])
        return [("", model["summary"]), ("Key timings. ", timings)]
    if section_key == "timeline":
        return [
            (f"{when}  ", f"{event} Source: {evidence}. Phase: {phase}.")
            for when, event, evidence, phase in model["timeline"]
        ]
    if section_key == "metrics":
        lines = [(f"{f['label']} ", f["caption"]) for f in model["figures"]]
        lines += [
            (
                "Metric summary row. ",
                f"{m}: {meaning}; normal {baseline}; during the incident {during}; source {source}.",
            )
            for m, meaning, baseline, during, source in model["metric_summary"]
        ]
        return lines
    if section_key == "customer_impact":
        return [(f"{question} ", answer) for question, answer in model["impact"]]
    if section_key == "response_analysis":
        lines = [("Went well: ", text) for text in model["went_well"]]
        lines += [(f"{area}: ", f"{gap} Effect: {effect}") for area, gap, effect in model["missed"]]
        lines += [("Channels of inquiry: ", "; ".join(model["channels"]))]
        return lines
    if section_key == "five_whys":
        lines = []
        for chain in model["chains"]:
            lines.append((f"Chain {chain['letter']}, {chain['title']}. Problem: ", chain["problem"]))
            lines += [(f"{label} ", text) for label, text in chain["whys"]]
            lines.append((f"RC-{chain['letter']}: ", chain["root_cause"]))
        return lines
    if section_key == "lessons":
        return [
            (f"{lid} ({rc}): ", f"{statement} Corrective actions: {cas}.")
            for lid, statement, rc, cas in model["lessons"]
        ]
    if section_key == "corrective_actions":
        return [
            (f"{cid} ({addresses}, improves {improves}): ", f"{statement} Action items {items}.")
            for cid, statement, addresses, improves, items in model["corrective_actions"]
        ]
    if section_key == "action_items":
        return [
            (f"{aid} ({ca}): ", f"{statement} Owner: {owner}. Due {due}. Status: {status}. {prevents}")
            for aid, statement, ca, owner, due, status, prevents in model["action_items"]
        ]
    if section_key == "review":
        lines = [
            (f"{stage}: ", f"{reviewer}, {date}. {outcome}")
            for stage, reviewer, date, outcome in model["review_stages"]
        ]
        lines += [
            ("Open question: ", f"{question} Owner: {owner}. Answer due {due}. Answer: {answer}.")
            for question, owner, due, answer in model["open_questions"]
        ]
        return lines
    raise ValueError(f"no example rendering for section {section_key}")


# ============================================================================ docx writing


def _shade(element: Any, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    if element.tag == qn("w:pPr"):
        _insert_in_order(element, shd, PPR_AFTER_BORDER[1:])
    else:
        element.append(shd)


#: CT_PPr children that must follow w:pBdr and w:shd, in schema order; used to insert in place.
PPR_AFTER_BORDER = tuple(
    qn(f"w:{tag}")
    for tag in (
        "shd",
        "tabs",
        "suppressAutoHyphens",
        "kinsoku",
        "wordWrap",
        "overflowPunct",
        "topLinePunct",
        "autoSpaceDE",
        "autoSpaceDN",
        "bidi",
        "adjustRightInd",
        "snapToGrid",
        "spacing",
        "ind",
        "contextualSpacing",
        "mirrorIndents",
        "suppressOverlap",
        "jc",
        "textDirection",
        "textAlignment",
        "textboxTightWrap",
        "outlineLvl",
        "divId",
        "cnfStyle",
        "rPr",
        "sectPr",
        "pPrChange",
    )
)
#: CT_Settings children that must follow w:updateFields, in schema order.
SETTINGS_AFTER_UPDATE_FIELDS = (
    *tuple(
        qn(f"w:{tag}")
        for tag in (
            "hdrShapeDefaults",
            "footnotePr",
            "endnotePr",
            "compat",
            "docVars",
            "rsids",
            "attachedSchema",
            "themeFontLang",
            "clrSchemeMapping",
            "doNotIncludeSubdocsInStats",
            "doNotAutoCompressPictures",
            "forceUpgrade",
            "captions",
            "readModeInkLockDown",
            "smartTagType",
            "schemaLibrary",
            "shapeDefaults",
            "doNotEmbedSmartTags",
            "decimalSymbol",
            "listSeparator",
        )
    ),
    qn("m:mathPr"),
)


def _insert_in_order(parent: Any, child: Any, successors: tuple[str, ...]) -> None:
    """Insert `child` before the first existing successor so the OOXML sequence stays valid."""
    for index, existing in enumerate(parent):
        if existing.tag in successors:
            parent.insert(index, child)
            return
    parent.append(child)


def _left_border(paragraph: Any, color: str) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "6")
    left.set(qn("w:color"), color)
    pbdr.append(left)
    _insert_in_order(ppr, pbdr, PPR_AFTER_BORDER)


def _field(paragraph: Any, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(end)


class DocxWriter:
    """Writes the layout into a python-docx Document."""

    def __init__(self) -> None:
        self.doc = Document()
        self._styles()

    def _styles(self) -> None:
        section = self.doc.sections[0]
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, side, Cm(2.0))
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        for name, size, color in (
            ("Heading 1", 16, NAVY),
            ("Heading 2", 12.5, ACCENT),
            ("Heading 3", 11, NAVY),
        ):
            style = self.doc.styles[name]
            style.font.name = FONT
            style.font.size = Pt(size)
            style.font.bold = True
            style.font.color.rgb = color
            style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        zoom = self.doc.settings.element.find(qn("w:zoom"))
        if zoom is not None and zoom.get(qn("w:percent")) is None:
            zoom.set(qn("w:percent"), "100")  # python-docx's default template omits the required attribute
        core = self.doc.core_properties
        core.author = "PIR Assessor"
        core.last_modified_by = "PIR Assessor"
        core.created = FIXED_TIMESTAMP
        core.modified = FIXED_TIMESTAMP
        core.title = "Post-Incident Review"

    # ---- runs and paragraphs
    def run(
        self,
        paragraph: Any,
        content: Cell,
        bold: bool = False,
        size: float | None = None,
        color: RGBColor | None = None,
    ) -> None:
        if isinstance(content, Placeholder):
            run = paragraph.add_run(f"[{content.text}]")
            run.italic = True
            run.font.color.rgb = PLACEHOLDER_TXT
            _shade(run._r.get_or_add_rPr(), PH_FILL)
        else:
            run = paragraph.add_run(content)
            run.bold = bold
            if color is not None:
                run.font.color.rgb = color
        if size is not None:
            run.font.size = Pt(size)

    def paragraph(self, *parts: Cell | tuple[Cell, bool]) -> Any:
        paragraph = self.doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(6)
        for part in parts:
            content, bold = part if isinstance(part, tuple) else (part, False)
            self.run(paragraph, content, bold=bold)
        return paragraph

    def heading(self, text: str, level: int) -> None:
        self.doc.add_heading(text, level=level)

    def guidance(self, text: str, label: str = "Guidance") -> None:
        paragraph = self.doc.add_paragraph()
        _shade(paragraph._p.get_or_add_pPr(), GUIDE_FILL)
        _left_border(paragraph, "7F7F7F")
        paragraph.paragraph_format.left_indent = Cm(0.2)
        paragraph.paragraph_format.space_after = Pt(8)
        self.run(paragraph, f"{label}: ", bold=True, size=9, color=GREY_TXT)
        self.run(paragraph, text, size=9, color=GREY_TXT)

    def example(self, lines: list[tuple[str, str]], label: str = "Example") -> None:
        for index, (lead, text) in enumerate(lines):
            paragraph = self.doc.add_paragraph()
            _shade(paragraph._p.get_or_add_pPr(), EX_FILL)
            _left_border(paragraph, "2E7D4F")
            paragraph.paragraph_format.left_indent = Cm(0.2)
            paragraph.paragraph_format.space_after = Pt(8 if index == len(lines) - 1 else 2)
            if index == 0:
                self.run(paragraph, f"{label}: ", bold=True, size=9.5, color=EXAMPLE_LABEL)
            if lead:
                self.run(paragraph, lead, bold=True, size=9.5, color=EXAMPLE_TXT)
            self.run(paragraph, text, size=9.5, color=EXAMPLE_TXT)

    def bullets(self, items: list[Cell], style: str = "List Bullet") -> None:
        for item in items:
            paragraph = self.doc.add_paragraph(style=style)
            self.run(paragraph, item)

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- tables
    def table(self, headers: list[str] | None, rows: list[Row], widths_cm: list[float]) -> None:
        assert abs(sum(widths_cm) - CONTENT_WIDTH_CM) < 0.01, (
            f"widths {widths_cm} do not sum to {CONTENT_WIDTH_CM}"
        )
        table = self.doc.add_table(rows=0, cols=len(widths_cm))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        table.autofit = False
        if headers:
            cells = table.add_row().cells
            for cell, header, width in zip(cells, headers, widths_cm, strict=True):
                cell.width = Cm(width)
                _shade(cell._tc.get_or_add_tcPr(), HEAD_FILL)
                cell.paragraphs[0].text = ""
                self.run(cell.paragraphs[0], header, bold=True, size=9.5, color=NAVY)
        for row in rows:
            cells = table.add_row().cells
            for cell, content, width in zip(cells, row, widths_cm, strict=True):
                cell.width = Cm(width)
                cell.paragraphs[0].text = ""
                self.run(cell.paragraphs[0], content, size=9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def kv_table(self, pairs: list[tuple[str, Cell]]) -> None:
        table = self.doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.autofit = False
        for key, value in pairs:
            cells = table.add_row().cells
            cells[0].width, cells[1].width = Cm(5.3), Cm(11.7)
            _shade(cells[0]._tc.get_or_add_tcPr(), HEAD_FILL)
            cells[0].paragraphs[0].text = ""
            self.run(cells[0].paragraphs[0], key, bold=True, size=9.5, color=NAVY)
            cells[1].paragraphs[0].text = ""
            self.run(cells[1].paragraphs[0], value, size=9.5)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # ---- document furniture
    def toc(self) -> None:
        paragraph = self.doc.add_paragraph()
        _field(paragraph, 'TOC \\o "1-2" \\h \\z \\u')
        settings = self.doc.settings.element
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        _insert_in_order(settings, update, SETTINGS_AFTER_UPDATE_FIELDS)

    def header_footer(self, incident_id: Cell, classification: Cell, label: str) -> None:
        section = self.doc.sections[0]
        header = section.header.paragraphs[0]
        header.text = ""
        self.run(header, "Post-Incident Review  ", size=8, color=GREY_TXT)
        self.run(header, incident_id, size=8)
        self.run(header, f"\t{label}", size=8, color=GREY_TXT)
        footer = section.footer.paragraphs[0]
        footer.text = ""
        self.run(footer, "Classification: ", size=8, color=GREY_TXT)
        self.run(footer, classification, size=8)
        self.run(footer, "\tPage ", size=8, color=GREY_TXT)
        _field(footer, "PAGE")
        self.run(footer, " of ", size=8, color=GREY_TXT)
        _field(footer, "NUMPAGES")

    def save(self, path: Path) -> None:
        """Write the document with every zip entry stamped `ZIP_DATE_TIME`, in python-docx's order."""
        written = io.BytesIO()
        self.doc.save(written)
        with zipfile.ZipFile(written) as source, zipfile.ZipFile(path, "w") as document:
            for info in source.infolist():
                entry = zipfile.ZipInfo(info.filename, date_time=ZIP_DATE_TIME)
                entry.compress_type = zipfile.ZIP_DEFLATED
                document.writestr(entry, source.read(info))


# ========================================================================= markdown writing


def _md_escape(content: Cell) -> str:
    text = f"[{content.text}]" if isinstance(content, Placeholder) else str(content)
    return text.replace("|", "\\|").replace("\n", " ")


class MarkdownWriter:
    """Writes the same layout as markdown; the structure, not the styling."""

    def __init__(self) -> None:
        self.lines: list[str] = []

    def run_text(self, content: Cell, bold: bool = False) -> str:
        text = f"[{content.text}]" if isinstance(content, Placeholder) else str(content)
        if bold and text.strip():
            return f"**{text.strip()}**" + (" " if text.endswith(" ") else "")
        return text

    def paragraph(self, *parts: Cell | tuple[Cell, bool]) -> None:
        pieces = [self.run_text(*(part if isinstance(part, tuple) else (part, False))) for part in parts]
        self.lines += ["".join(pieces), ""]

    def heading(self, text: str, level: int) -> None:
        self.lines += [f"{'#' * level} {text}", ""]

    def guidance(self, text: str, label: str = "Guidance") -> None:
        self.lines += [f"> **{label}:** {text}", ""]

    def example(self, lines: list[tuple[str, str]], label: str = "Example") -> None:
        for index, (lead, text) in enumerate(lines):
            prefix = f"**{label}:** " if index == 0 else ""
            lead_md = f"**{lead.strip()}** " if lead.strip() else ""
            self.lines.append(f"> {prefix}{lead_md}{text}")
            if index < len(lines) - 1:
                self.lines.append(">")
        self.lines.append("")

    def bullets(self, items: list[Cell], style: str = "List Bullet") -> None:
        for index, item in enumerate(items, start=1):
            marker = f"{index}." if style == "List Number" else "-"
            self.lines.append(f"{marker} {self.run_text(item)}")
        self.lines.append("")

    def page_break(self) -> None:
        self.lines += ["---", ""]

    def table(self, headers: list[str] | None, rows: list[Row], widths_cm: list[float]) -> None:
        columns = len(widths_cm)
        head = headers or [""] * columns
        self.lines.append("| " + " | ".join(head) + " |")
        self.lines.append("|" + " --- |" * columns)
        for row in rows:
            self.lines.append("| " + " | ".join(_md_escape(c) for c in row) + " |")
        self.lines.append("")

    def kv_table(self, pairs: list[tuple[str, Cell]]) -> None:
        self.table(["Field", "Value"], [[k, v] for k, v in pairs], [5.3, 11.7])

    def toc(self) -> None:
        self.lines += ["_Contents: generated by the document tool from the headings below._", ""]

    def header_footer(self, incident_id: Cell, classification: Cell, label: str) -> None:
        self.lines += [f"_{label}. Classification: {self.run_text(classification)}._", ""]

    def save(self, path: Path) -> None:
        path.write_text(self.text(), encoding="utf-8")

    def text(self) -> str:
        return "\n".join(self.lines).rstrip() + "\n"


# ============================================================================ the layout


def render(
    writer: DocxWriter | MarkdownWriter,
    model: dict[str, Any],
    sections: dict[str, dict[str, Any]],
    principles: dict[str, Any],
    is_template: bool,
    appendix: bool,
    example_model: dict[str, Any] | None,
) -> None:
    label = f"TEMPLATE {TEMPLATE_VERSION}" if is_template else f"{model['incident_id']}"
    writer.header_footer(model["incident_id"], model["classification"], label)

    def guide(key: str) -> None:
        if is_template:
            writer.guidance(str(sections[key]["template_guidance"]))

    def title(key: str, number: int) -> str:
        return f"{number}. {BODY_TITLE_OVERRIDES.get(key, sections[key]['title'])}"

    # Title block
    writer.paragraph(("Post-Incident Review (PIR)", True))
    writer.paragraph(model["incident_id"], "  ", model["title"])
    writer.paragraph("Blameless review. Processes, not people.")
    if is_template:
        writer.guidance(str(sections["cross_cutting"]["template_guidance"]), "How to use this template")
        writer.guidance(
            "Write the sections in the order the reference guidance recommends: 2 Timeline, 3 Metrics, 4 Customer "
            "impact, 5 Response analysis, 6 Five Whys, 7 Lessons, 8 Corrective actions, 9 Action items, and write "
            "1 Summary last. The document is laid out in reading order, not writing order. Appendix A explains what "
            "good looks like for each section with one worked example that runs through every section; Appendix B "
            "summarises PIR principles, timing, presenting and reviewing. Remove the guidance boxes and the "
            "appendices before publishing. Start the PIR as soon as customer impact is mitigated; draft within 2 to "
            "3 business days; publish within 1 to 2 weeks.",
            "Writing order",
        )

    writer.heading("Document control", 1)
    guide("document_control")
    writer.kv_table(model["control"])

    writer.heading("Contents", 1)
    writer.toc()
    if is_template:
        writer.guidance(
            "Right-click and choose Update Field (or press F9) in Word to populate the contents after editing."
        )
    writer.page_break()

    # 1 Summary
    writer.heading(title("summary", 1), 1)
    guide("summary")
    writer.paragraph(model["summary"])
    writer.heading("1.1 Key timings", 2)
    if is_template:
        writer.guidance(
            'Derived from the timeline in section 2. Report literal values with units; do not estimate. If a value is unknown, say "not determined" and why.'
        )
    writer.table(["Measure", "Definition", "Value"], [list(r) for r in model["key_timings"]], [5.0, 7.8, 4.2])

    # 2 Timeline
    writer.heading(title("timeline", 2), 1)
    guide("timeline")
    writer.table(
        ["Time (TZ)", "Event (what happened, what was observed or decided)", "Evidence / source", "Phase"],
        model["timeline"] or [[NOT_PROVIDED, "", "", ""]],
        [3.0, 8.4, 3.3, 2.3],
    )

    # 3 Metrics
    writer.heading(title("metrics", 3), 1)
    guide("metrics")
    writer.heading("3.1 Figures", 2)
    if not model["figures"]:
        writer.paragraph(NOT_PROVIDED)
    for figure in model["figures"]:
        if figure["insert"] is not None:
            writer.paragraph(figure["insert"])
        writer.paragraph((figure["label"] + " ", True), figure["caption"])
    writer.heading("3.2 Metric summary", 2)
    writer.table(
        ["Metric", "What it tells us", "Normal baseline", "During incident", "Source / link"],
        model["metric_summary"] or [[NOT_PROVIDED, "", "", "", ""]],
        [4.0, 3.8, 2.6, 2.6, 4.0],
    )

    # 4 Customer impact
    writer.heading(title("customer_impact", 4), 1)
    guide("customer_impact")
    writer.table(["Question", "Answer"], [[q, a] for q, a in model["impact"]], [6.3, 10.7])

    # 5 Response analysis
    writer.heading(title("response_analysis", 5), 1)
    guide("response_analysis")
    writer.heading("5.1 What went well", 2)
    writer.bullets(model["went_well"] or [NOT_PROVIDED])
    writer.heading("5.2 What we missed", 2)
    writer.table(
        ["Area", "What we missed or what slowed us down", "Effect on the incident"],
        model["missed"] or [[NOT_PROVIDED, "", ""]],
        [4.2, 8.0, 4.8],
    )
    writer.heading("5.3 Channels of inquiry", 2)
    if is_template:
        writer.guidance(
            "List each distinct problem that contributed. Each becomes a Five Whys chain in section 6."
        )
    writer.bullets(model["channels"] or [NOT_PROVIDED], style="List Number")

    # 6 Five Whys
    writer.heading(title("five_whys", 6), 1)
    guide("five_whys")
    if not model["chains"]:
        writer.paragraph(NOT_PROVIDED)
    for index, chain in enumerate(model["chains"], start=1):
        writer.heading(f"6.{index} Chain {chain['letter']}: {chain['title']}", 2)
        writer.paragraph(("Problem statement: ", True), chain["problem"])
        writer.table(
            ["Why?", "Because (answer, with evidence)"],
            [[label, text] for label, text in chain["whys"]],
            [5.5, 11.5],
        )
        writer.paragraph((f"Root cause RC-{chain['letter']}: ", True), chain["root_cause"])

    # 7 Lessons
    writer.heading(title("lessons", 7), 1)
    guide("lessons")
    writer.table(
        ["ID", "Lesson (what should have happened differently)", "From root cause", "Corrective action(s)"],
        model["lessons"] or [[NOT_PROVIDED, "", "", ""]],
        [1.6, 9.4, 2.8, 3.2],
    )

    # 8 Corrective actions
    writer.heading(title("corrective_actions", 8), 1)
    guide("corrective_actions")
    writer.table(
        ["ID", "Corrective action (what we will change)", "Addresses", "Primarily improves", "Action items"],
        model["corrective_actions"] or [[NOT_PROVIDED, "", "", "", ""]],
        [1.6, 7.6, 2.1, 3.0, 2.7],
    )

    # 9 Action items
    writer.heading(title("action_items", 9), 1)
    guide("action_items")
    writer.table(
        [
            "ID",
            "Action item (specific, verifiable)",
            "CA",
            "Owner (agreed)",
            "Due date",
            "Status",
            "How it prevents recurrence",
        ],
        model["action_items"] or [[NOT_PROVIDED, "", "", "", "", "", ""]],
        [1.4, 5.0, 1.2, 2.7, 1.9, 1.6, 3.2],
    )

    # 10 Review
    writer.heading(title("review", 10), 1)
    guide("review")
    if is_template:
        writer.guidance(
            'Reviewer of this template: confirm the expansion of "GO" used by your organisation; it is carried over from the reference guidance.',
            "Note",
        )
    writer.table(
        ["Review stage", "Reviewer (role)", "Date", "Outcome and material changes requested"],
        model["review_stages"] or [[NOT_PROVIDED, "", "", ""]],
        [4.6, 3.9, 2.5, 6.0],
    )
    writer.heading("10.1 Open questions raised in review", 2)
    if is_template:
        writer.guidance(
            'It is fine to say "I don\'t know, I will get back to you". Record the question, who owns the answer, and when.'
        )
    writer.table(
        ["Question", "Owner", "Answer due", "Answer / resolution"],
        model["open_questions"] or [["none", "", "", ""]],
        [6.3, 3.2, 2.5, 5.0],
    )

    if not appendix:
        return

    # Appendix A
    writer.page_break()
    writer.heading("Appendix A. Section guidance and worked example", 1)
    writer.paragraph(
        "This appendix states the expectation for each section of the template and shows what good looks like. "
        "The examples all describe one fictional incident, rendered from the same packet the validator accepts, so "
        "that you can see how the sections connect: a lesson references a root cause, a corrective action references "
        "a lesson, and an action item references a corrective action."
    )
    if example_model is not None:
        writer.example(
            [
                (
                    "Fictional incident used throughout: ",
                    f"{example_model['incident_id']}, {example_model['title']}. The organisation, services, numbers and people are invented.",
                )
            ],
            "Worked example",
        )
    for number, key in enumerate(BODY_ORDER, start=1):
        entry = sections[key]
        writer.heading(f"A.{number} {entry['title']}", 2)
        writer.paragraph(("Expectation. ", True), str(entry["expectation"]).strip())
        mistakes = " ".join(str(m).strip().rstrip(".") + "." for m in entry.get("common_mistakes") or [])
        writer.paragraph(("Common mistakes. ", True), mistakes)
        if example_model is not None:
            writer.example(example_lines(key, example_model))

    # Appendix B
    writer.page_break()
    writer.heading("Appendix B. PIR principles, timing, presenting and reviewing", 1)
    writer.paragraph(str(principles["preamble"]).strip())
    for number, section in enumerate(principles["sections"], start=1):
        writer.heading(f"B.{number} {section['title']}", 2)
        writer.bullets([str(point) for point in section["points"]])


# ================================================================================= the CLI


def load_principles(path: Path = PRINCIPLES_PATH) -> dict[str, Any]:
    data = yaml.safe_load(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("sections"), list):
        raise ValueError(f"{path}: not a principles catalog")
    return data


def build(
    packet_path: Path | None,
    output_format: str,
    appendix: bool,
    example_packet_path: Path = DEFAULT_EXAMPLE_PACKET,
) -> DocxWriter | MarkdownWriter:
    sections = pir_interview.load_section_catalog()
    principles = load_principles()
    if packet_path is None:
        model = template_model()
        example_model: dict[str, Any] | None = packet_model(validate_pir.load_document(example_packet_path))
    else:
        model = packet_model(validate_pir.load_document(packet_path))
        example_model = packet_model(validate_pir.load_document(example_packet_path)) if appendix else None
    writer: DocxWriter | MarkdownWriter = DocxWriter() if output_format == "docx" else MarkdownWriter()
    render(
        writer,
        model,
        sections,
        principles,
        is_template=packet_path is None,
        appendix=appendix,
        example_model=example_model,
    )
    return writer


def comparable_text(writer: DocxWriter | MarkdownWriter, scratch: Path) -> str:
    """What --check compares: the markdown itself, or the docx's outline."""
    if isinstance(writer, MarkdownWriter):
        return writer.text()
    writer.save(scratch)
    return docx_outline.outline(scratch)


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Render the PIR template, or a PIR packet, as .docx or markdown."
    )
    parser.add_argument(
        "--packet", help="A PIR packet to render. Without it, the blank template is rendered."
    )
    parser.add_argument("--format", choices=("docx", "md"), default="docx")
    parser.add_argument(
        "--appendix",
        dest="appendix",
        action="store_true",
        help="Include Appendices A and B (default for the template).",
    )
    parser.add_argument(
        "--no-appendix",
        dest="appendix",
        action="store_false",
        help="Omit the appendices (default for a packet).",
    )
    parser.set_defaults(appendix=None)
    parser.add_argument(
        "--example-packet",
        default=str(DEFAULT_EXAMPLE_PACKET),
        help="The packet Appendix A's worked example is rendered from.",
    )
    parser.add_argument(
        "--output", help="Write here. Required for docx unless --check; markdown defaults to stdout."
    )
    parser.add_argument(
        "--check",
        help="Compare with this expected rendering (markdown text, or a docx outline) and exit 1 with a diff if they differ.",
    )
    args = parser.parse_args(argv)
    if args.appendix is None:
        args.appendix = args.packet is None
    if args.output and args.check:
        parser.error("--output and --check are exclusive")
    if args.format == "docx" and not args.output and not args.check:
        parser.error("--format docx needs --output (or --check)")
    return args


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    packet_path = Path(args.packet) if args.packet else None
    try:
        writer = build(packet_path, args.format, args.appendix, Path(args.example_packet))
    except FileNotFoundError as exc:
        print(f"ERROR: file not found: {exc.filename}", file=sys.stderr)
        return 1
    except (OSError, ValueError, yaml.YAMLError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    if args.check:
        expected_path = Path(args.check)
        if not expected_path.is_file():
            print(f"ERROR: --check target {expected_path} does not exist", file=sys.stderr)
            return 1
        # The comparison needs a real .docx on disk; it is written to a private temporary directory,
        # never beside the expected file, so two concurrent checks cannot race and nothing is left behind.
        with tempfile.TemporaryDirectory(prefix="pir-docx-check-") as scratch_dir:
            current = comparable_text(writer, Path(scratch_dir) / expected_path.with_suffix(".docx").name)
        expected = expected_path.read_text(encoding="utf-8")
        if current == expected:
            print(f"PASS {expected_path} is current", file=sys.stderr)
            return 0
        sys.stderr.write(
            "".join(
                difflib.unified_diff(
                    expected.splitlines(keepends=True),
                    current.splitlines(keepends=True),
                    fromfile=str(expected_path),
                    tofile="re-rendered",
                )
            )
        )
        print(f"FAIL {expected_path} is stale; re-render it with --output", file=sys.stderr)
        return 1

    if args.output:
        writer.save(Path(args.output))
        print(f"wrote {args.output}", file=sys.stderr)
    elif isinstance(writer, MarkdownWriter):
        sys.stdout.write(writer.text())
    else:
        # parse_args refuses `--format docx` without --output, so only markdown reaches stdout.
        raise TypeError("a .docx has no text form; it is written with --output")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
